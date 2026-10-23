"""Compara dos carpetas de salidas (output_api parquet + Excel Anexo D) por contenido.

Uso: python comparar_salidas.py <ref_api_dir> <ref_xlsx> <nuevo_api_dir> <nuevo_xlsx>
"""
import sys
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import pyarrow.parquet as pq

ref_api, ref_x, new_api, new_x = map(Path, sys.argv[1:5])

print("== parquet")
for f in sorted(ref_api.glob("*.parquet")):
    g = new_api / f.name
    a, b = pd.read_parquet(f), pd.read_parquet(g)
    sa, sb = pq.read_schema(f), pq.read_schema(g)
    msg = []
    if list(a.columns) != list(b.columns):
        msg.append(f"columnas distintas {set(a.columns) ^ set(b.columns)}")
    if len(a) != len(b):
        msg.append(f"filas {len(a)} vs {len(b)}")
    if not sa.equals(sb, check_metadata=False):
        msg.append("esquema arrow distinto")
    if not msg:
        try:
            pd.testing.assert_frame_equal(a, b, check_exact=True)
            msg.append("IDENTICO (valores exactos, mismo orden)")
        except AssertionError:
            num = a.select_dtypes("number").columns
            dif = {c: float(np.nanmax(np.abs(a[c].to_numpy(float) - b[c].to_numpy(float))))
                   for c in num if not np.array_equal(a[c].to_numpy(), b[c].to_numpy(), equal_nan=True)}
            obj = [c for c in a.columns if c not in num and not a[c].equals(b[c])]
            msg.append(f"VALORES DISTINTOS num_maxabs={dif} no_num={obj}")
            for c, d in dif.items():
                rel = float(np.nansum(np.abs(a[c] - b[c])) / max(np.nansum(np.abs(a[c])), 1e-12))
                ncel = int((~np.isclose(a[c].to_numpy(float), b[c].to_numpy(float), rtol=0, atol=0, equal_nan=True)).sum())
                msg.append(f"   {c}: celdas_dif={ncel} dif_rel_total={rel:.3e}")
    print(f.name, "|", "; ".join(msg))

print("== xlsx (zip)")
za, zb = zipfile.ZipFile(ref_x), zipfile.ZipFile(new_x)
na, nb = set(za.namelist()), set(zb.namelist())
print("miembros solo ref:", na - nb, "solo nuevo:", nb - na)
for n in sorted(na & nb):
    if za.read(n) != zb.read(n):
        print("  miembro distinto:", n)
print("== xlsx (celdas)")
wa = openpyxl.load_workbook(ref_x, data_only=False)
wb = openpyxl.load_workbook(new_x, data_only=False)
print("hojas", wa.sheetnames == wb.sheetnames, wa.sheetnames)
tot = 0
for s in wa.sheetnames:
    A, B = wa[s], wb[s]
    n = 0
    ej = []
    for r in range(1, max(A.max_row, B.max_row) + 1):
        for c in range(1, max(A.max_column, B.max_column) + 1):
            va, vb = A.cell(r, c).value, B.cell(r, c).value
            if va != vb:
                n += 1
                if len(ej) < 5:
                    ej.append((A.cell(r, c).coordinate, va, vb))
    tot += n
    print(f"  {s}: dims {A.max_row}x{A.max_column} vs {B.max_row}x{B.max_column}; celdas distintas {n}; ej {ej}")
print("TOTAL celdas distintas:", tot)
