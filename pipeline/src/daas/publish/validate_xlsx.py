"""Validacion del Anexo D ya generado.

Abre el `.xlsx` de entrega y comprueba lo que el banco va a mirar primero, sin releer los contratos:

1. Hojas obligatorias y encabezados oficiales (fila 2).
2. Listas de validacion de la plantilla: cada valor escrito pertenece a su lista y el rango de la lista
   cubre todas las filas escritas (la plantilla trae los rangos hasta ~la fila 1013).
3. Columnas que deben ser enteras (`validacion.enteras`; se permite "No aplica" y las formulas).
4. Columnas obligatorias no vacias (`validacion.obligatorias`).
5. RN-02: la hoja 01 trae cada (modelo, entidad, objetivo) con cada dia del horizonte, sin repetir, y toda
   entidad M1-M3 de la hoja 01 tiene al menos una fila en la hoja 02.
6. Las filas `sin_accion` de la hoja 02 llevan monto 0 y "No aplica" en saldo y costo.

Los umbrales, columnas y vocabulario salen de ``conf/export_anexo_d.yml``; nada se escribe aqui.
Una lectura del archivo, sin tocar S3 ni los contratos.
"""

from __future__ import annotations

import datetime as _dt
import re
from pathlib import Path
from typing import Any

#: Maximo de ejemplos que se citan por incumplimiento (el conteo siempre es completo).
_EXAMPLES = 5


def _literal_list(formula: str | None) -> list[str] | None:
    """Valores de una lista de validacion literal (``"a,b,c"``); ``None`` si es un rango o formula."""
    if not formula:
        return None
    text = str(formula).strip()
    if len(text) < 2 or text[0] != '"' or text[-1] != '"':
        return None
    return [v.strip() for v in text[1:-1].split(",")]


def _to_date_or_none(value: object) -> _dt.date | None:
    """Fecha de una celda: `date`/`datetime` o el numero de serie de Excel (celda sin formato de fecha)."""
    if isinstance(value, _dt.datetime):
        return value.date()
    if isinstance(value, _dt.date):
        return value
    if isinstance(value, int | float) and not isinstance(value, bool):
        from openpyxl.utils.datetime import from_excel

        return from_excel(value).date()
    return None


def _cite(build_examples: list[Any], total: int) -> str:
    """`total` incumplimientos con los primeros ejemplos."""
    sample = ", ".join(repr(e) for e in build_examples[:_EXAMPLES])
    return f"{total} filas (ej.: {sample})"


def _columns(ws: Any) -> dict[str, int]:
    """Encabezado (fila 2) -> indice de columna (1-based)."""
    return {str(c.value): c.column for c in ws[2] if c.value is not None}


def _data_rows(ws: Any) -> list[int]:
    """Filas de datos escritas (fila 3 en adelante, con id_proveedor en la columna A)."""
    return [r for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=1).value not in (None, "")]


def _validate_lists(ws: Any, sheet: str, rows: list[int], errors: list[str]) -> None:
    last = max(rows) if rows else 2
    header = {c: n for n, c in _columns(ws).items()}
    for dv in ws.data_validations.dataValidation:
        if dv.type != "list":
            continue
        allowed = _literal_list(dv.formula1)
        if allowed is None:
            continue
        for rng in dv.sqref.ranges:
            for col in range(rng.min_col, rng.max_col + 1):
                name = header.get(col, f"col {col}")
                if rng.max_row < last:
                    errors.append(
                        f"{sheet}.{name}: la lista de validacion llega a la fila {rng.max_row} y hay "
                        f"datos hasta la {last}."
                    )
                invalid_values = []
                for r in rows:
                    if not (rng.min_row <= r <= rng.max_row):
                        continue
                    v = ws.cell(row=r, column=col).value
                    if v in (None, "") or (isinstance(v, str) and v.startswith("=")):
                        continue
                    if str(v) not in allowed:
                        invalid_values.append(v)
                if invalid_values:
                    errors.append(
                        f"{sheet}.{name}: valores fuera de la lista {allowed}: "
                        f"{_cite(invalid_values, len(invalid_values))}."
                    )


def _validate_integer_and_required(
    ws: Any, sheet: str, rows: list[int], cfg: dict[str, Any], errors: list[str]
) -> None:
    columns = _columns(ws)
    val = cfg["validacion"]
    not_applicable = cfg["vocabulario"]["no_aplica"]
    for name in (val.get("enteras") or {}).get(sheet, []):
        col = columns.get(name)
        if col is None:
            errors.append(f"{sheet}: falta la columna entera {name!r}.")
            continue
        invalid_values = []
        for r in rows:
            v = ws.cell(row=r, column=col).value
            if v in (None, "", not_applicable) or (isinstance(v, str) and v.startswith("=")):
                continue
            if isinstance(v, bool) or not isinstance(v, int | float) or float(v) != int(v):
                invalid_values.append(v)
        if invalid_values:
            errors.append(f"{sheet}.{name}: debe ser un entero en COP: {_cite(invalid_values, len(invalid_values))}.")
    for name in (val.get("obligatorias") or {}).get(sheet, []):
        col = columns.get(name)
        if col is None:
            errors.append(f"{sheet}: falta la columna obligatoria {name!r}.")
            continue
        empty = [r for r in rows if ws.cell(row=r, column=col).value in (None, "")]
        if empty:
            errors.append(f"{sheet}.{name}: obligatoria y vacia: {_cite(empty, len(empty))} (filas de Excel).")


def _validate_rn02(wb: Any, cfg: dict[str, Any], errors: list[str]) -> None:
    """RN-02: cobertura completa de la hoja 01 por dia del horizonte y entidad con plan en la hoja 02."""
    h1, h2 = wb["01_Predicciones"], wb["02_Plan_Operativo"]
    c1, c2 = _columns(h1), _columns(h2)
    start = _to_date_or_none(cfg["horizonte"]["inicio"]) or _dt.date.fromisoformat(str(cfg["horizonte"]["inicio"]))
    fin = _to_date_or_none(cfg["horizonte"]["fin"]) or _dt.date.fromisoformat(str(cfg["horizonte"]["fin"]))
    expected = {start + _dt.timedelta(days=i) for i in range((fin - start).days + 1)}
    by_series: dict[tuple[str, str, str], list[_dt.date | None]] = {}
    for r in _data_rows(h1):
        key = tuple(str(h1.cell(row=r, column=c1[n]).value) for n in ("modelo", "id_entidad", "variable_objetivo"))
        by_series.setdefault(key, []).append(_to_date_or_none(h1.cell(row=r, column=c1["fecha_objetivo"]).value))
    incomplete, repeated = [], []
    for key, dates in by_series.items():
        if len(dates) != len(set(dates)):
            repeated.append("/".join(key))
        if set(dates) != expected:
            incomplete.append("/".join(key))
    if incomplete:
        errors.append(
            f"RN-02 hoja 01: series sin los {len(expected)} dias del horizonte: {_cite(incomplete, len(incomplete))}."
        )
    if repeated:
        errors.append(f"RN-02 hoja 01: series con fechas repetidas: {_cite(repeated, len(repeated))}.")

    plan = {
        (str(h2.cell(row=r, column=c2["modelo"]).value), str(h2.cell(row=r, column=c2["id_entidad"]).value))
        for r in _data_rows(h2)
    }
    target = {(m, e) for m, e, _ in by_series if m in {"M1", "M2", "M3"}}
    without_plan = sorted(f"{m}/{e}" for m, e in target - plan)
    if without_plan:
        errors.append(
            f"RN-02 hoja 02: entidades de la hoja 01 sin ninguna fila de plan: "
            f"{_cite(without_plan, len(without_plan))}."
        )

    not_applicable = cfg["vocabulario"]["no_aplica"]
    invalid_rows = []
    for r in _data_rows(h2):
        if h2.cell(row=r, column=c2["accion_sugerida"]).value != "sin_accion":
            continue
        ok = (
            h2.cell(row=r, column=c2["monto_sugerido_cop"]).value == 0
            and h2.cell(row=r, column=c2["saldo_antes_accion_cop"]).value == not_applicable
            and h2.cell(row=r, column=c2["costo_estimado_cop"]).value == not_applicable
        )
        if not ok:
            invalid_rows.append(r)
    if invalid_rows:
        errors.append(
            f"Hoja 02: filas sin_accion sin monto 0 y '{not_applicable}' en saldo y costo: "
            f"{_cite(invalid_rows, len(invalid_rows))} (filas de Excel)."
        )


#: Limite de caracteres de una celda de Excel.
_MAX_CELL = 32767


def _sentences(text: str) -> list[str]:
    return [p.strip() for p in re.split(r"(?<=\.) ", text) if p.strip()]


def _validate_texts(wb: Any, cfg: dict[str, Any], errors: list[str]) -> None:
    """Textos para el evaluador (interruptor `hoja04_v2`; hoja 03 con `hoja03_textos_v2`).

    En la hoja 04 (todas las celdas de texto) y, si aplica, en `comentarios` y `metodologia_validacion` de la hoja 03:
    (a) terminos prohibidos de `hoja_04_v2.prohibidos`; (b) palabras de `tildes_control` sin tilde (palabra completa,
    sin distinguir mayusculas; no cuenta lo que va entre comillas invertidas ni unido por guion bajo); (c) oraciones
    repetidas en una celda; (d) SUP citado sin titulo en `sup_titulos` (hoja 04); (e) celda > 32.767 caracteres.
    """
    from daas import alx

    if not alx.interruptor("hoja04_v2"):
        return
    h4 = cfg["hoja_04_v2"]
    banned = [re.compile(p) for p in h4["prohibidos"]]
    words = re.compile(r"\b(" + "|".join(re.escape(w) for w in h4["tildes_control"]) + r")\b", re.IGNORECASE)
    prefixes = re.compile(r"\b(" + "|".join(re.escape(w) for w in h4.get("tildes_prefijos", [])) + r")", re.IGNORECASE)
    sup = re.compile(r"SUP-(M3-\d{2}|\d{2})")
    titles = {str(k) for k in h4["sup_titulos"]}
    targets: list[tuple[str, list[str] | None]] = [("04_Descripcion_Solucion", None)]
    if alx.interruptor("hoja03_textos_v2"):
        targets.append(("03_Metricas_Desarrollo", ["comentarios", "metodologia_validacion"]))
    for sheet, names in targets:
        ws = wb[sheet]
        columns = _columns(ws)
        cols = [columns[n] for n in names] if names else list(columns.values())
        problems: dict[str, list[str]] = {}
        for r in _data_rows(ws):
            for col in cols:
                v = ws.cell(row=r, column=col).value
                if not isinstance(v, str) or v.startswith("="):
                    continue
                where = f"fila {r} col {col}"
                if len(v) > _MAX_CELL:
                    problems.setdefault("celda de mas de 32.767 caracteres", []).append(where)
                for p in banned:
                    m = p.search(v)
                    if m:
                        problems.setdefault(f"termino prohibido {p.pattern!r}", []).append(f"{where}: {m.group(0)!r}")
                plain = re.sub(r"`[^`]*`", "", v)
                for m in list(words.finditer(plain)) + list(prefixes.finditer(plain)):
                    problems.setdefault("palabra sin tilde", []).append(f"{where}: {m.group(0)!r}")
                sentences = _sentences(v)
                if len(sentences) != len(set(sentences)):
                    problems.setdefault("oracion repetida", []).append(where)
                if sheet == "04_Descripcion_Solucion":
                    missing = sorted({c for c in sup.findall(v) if c not in titles})
                    if missing:
                        problems.setdefault("SUP sin titulo", []).append(f"{where}: {missing}")
        for kind, items in problems.items():
            errors.append(f"{sheet}: {kind}: {_cite(items, len(items))}.")


def validate_anexo_d(file_path: str | Path, cfg: dict[str, Any]) -> list[str]:
    """Valida el `.xlsx` del Anexo D contra la plantilla y la configuracion.

    Args:
        file_path: Ruta del `.xlsx` generado.
        cfg: Configuracion de ``export_anexo_d.yml`` (bloques `vocabulario`, `horizonte`, `validacion`).

    Returns:
        list[str]: Un mensaje por incumplimiento; vacia si el archivo cumple todo.
    """
    import openpyxl

    from daas.publish.anexo_d import HEADERS, REQUIRED_SHEETS

    wb = openpyxl.load_workbook(file_path, data_only=False)
    errors: list[str] = []
    missing = [h for h in REQUIRED_SHEETS if h not in wb.sheetnames]
    if missing:
        return [f"Faltan hojas obligatorias: {missing}."]
    for sheet in REQUIRED_SHEETS:
        ws = wb[sheet]
        expected = HEADERS[sheet]
        real = [ws.cell(row=2, column=c).value for c in range(1, len(expected) + 1)]
        if real != expected:
            errors.append(f"{sheet}: encabezados distintos de los oficiales: {real}.")
            continue
        rows = _data_rows(ws)
        if not rows:
            errors.append(f"{sheet}: sin filas de datos.")
            continue
        _validate_lists(ws, sheet, rows, errors)
        _validate_integer_and_required(ws, sheet, rows, cfg, errors)
    if not errors:
        _validate_rn02(wb, cfg, errors)
    _validate_texts(wb, cfg, errors)
    return errors
